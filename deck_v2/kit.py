"""Design kit for the DSASSASSIN v2 deck — Apple-inspired, no animations.

Principles applied: generous whitespace, strong typographic hierarchy, one idea per slide, a single
accent colour used sparingly, no decorative chrome (no header bars, borders, chevrons, progress bars),
hairline tables, large centred figures. Transitions: a quiet fade only.
"""
from pathlib import Path
import json

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
ASSET_DIRS = [ROOT / 'presentation_assets', ROOT / 'figs']
OUT = HERE / 'output'
OUT.mkdir(exist_ok=True)

# Apple-like palette
C = dict(black='1D1D1F', gray='6E6E73', gray3='86868B', hair='D2D2D7', fill='F5F5F7', fill2='E8E8ED',
         white='FFFFFF', accent='0071E3', red='FF3B30', redtext='D70015', green='248A3D', orange='C93400',
         bg_dark='000000', fill_dark='1C1C1E', gray_dark='A1A1A6', hair_dark='3A3A3C', num_dark='2C2C2E')
FONT = 'Segoe UI'
FONT_SEMI = 'Segoe UI Semibold'
SW, SH = 13.333333, 7.5
M = 0.9                    # side margin
CW = SW - 2 * M            # content width
TOP = 1.95                 # content top (below title)

P = Presentation()
P.slide_width = Inches(SW)
P.slide_height = Inches(SH)
P.core_properties.title = 'DSASSASSIN — Cross-VM side channels through Intel DSA'
P.core_properties.author = 'Paper presentation | Chen et al., HPCA 2026'

MANIFEST = []
_UID = [0]
_DARK = {}


def rgb(c):
    return RGBColor.from_string(C.get(c, c))


def uid(prefix='sh'):
    _UID[0] += 1
    return f'{prefix}_{_UID[0]:04d}'


def notes(text):
    return text


def is_dark(s):
    return _DARK.get(id(s), False)


# ----------------------------------------------------------------------------- primitives
def rect(parent, x, y, w, h, fill='fill', line=None, radius=0.0, shape=None, line_w=0.75, dash=False):
    typ = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = parent.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(line_w)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sh.line.fill.background()
    if radius and typ == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.shadow.inherit = False
    sh.name = uid('rect')
    return sh


def hline(parent, x, y, w, color='hair', width=0.75):
    ln = parent.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(width)
    ln.name = uid('hl')
    return ln


def arrow(parent, x1, y1, x2, y2, color='gray3', width=1.25):
    ln = parent.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(width)
    lnxml = ln.line._get_or_add_ln()
    tail = etree.SubElement(lnxml, qn('a:tailEnd')); tail.set('type', 'triangle'); tail.set('w', 'sm'); tail.set('len', 'med')
    ln.name = uid('ar')
    return ln


def _para(par, text, size, color, font, bold=False, italic=False, align=None, space_after=0, line_spacing=None):
    par.text = text
    par.font.name = font; par.font.size = Pt(size); par.font.bold = bold; par.font.italic = italic
    par.font.color.rgb = rgb(color)
    par.space_after = Pt(space_after)
    if line_spacing:
        par.line_spacing = line_spacing
    if align is not None:
        par.alignment = align


def tx(parent, text, x, y, w, h, size=20, color='black', semi=False, bold=False, align=None, anchor=None,
       space_after=4, line_spacing=None, italic=False, font=None):
    sh = parent.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    if anchor is not None:
        tf.vertical_anchor = anchor
    f = font or (FONT_SEMI if semi else FONT)
    for i, ln in enumerate(str(text).split('\n')):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _para(par, ln, size, color, f, bold, italic, align, space_after, line_spacing)
    sh.name = uid('tx')
    return sh


def chip(parent, text, x, y, w, h, fill='fill', color='black', size=15, semi=True, radius=.16, line=None, dash=False):
    sh = rect(parent, x, y, w, h, fill, line=line, radius=radius, dash=dash)
    tf = sh.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.08); tf.margin_top = tf.margin_bottom = Inches(.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(str(text).split('\n')):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _para(par, ln, size, color, FONT_SEMI if semi else FONT, align=PP_ALIGN.CENTER)
    return sh


# ----------------------------------------------------------------------------- transitions
_NS_MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
_NS_P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
_NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def set_fade(slide_, dur_ms=500):
    xml = (f'<mc:AlternateContent xmlns:mc="{_NS_MC}" xmlns:p="{_NS_P}" xmlns:p14="{_NS_P14}">'
           f'<mc:Choice Requires="p14"><p:transition spd="med" p14:dur="{dur_ms}"><p:fade/></p:transition></mc:Choice>'
           f'<mc:Fallback><p:transition spd="med"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    el = etree.fromstring(xml)
    sld = slide_._element
    anchor = sld.find(qn('p:clrMapOvr'))
    if anchor is None:
        anchor = sld.find(qn('p:cSld'))
    anchor.addnext(el)


# ----------------------------------------------------------------------------- slide frame
def slide(title, act, source, notes_text, secs=75, dark=False, backup=False, title_size=30):
    s = P.slides.add_slide(P.slide_layouts[6])
    _DARK[id(s)] = dark
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb('bg_dark' if dark else 'white')
    set_fade(s)
    fg, sec = ('white', 'gray_dark') if dark else ('black', 'gray3')
    if act:
        tx(s, ('Backup · ' if backup else '') + act, M, .55, CW, .3, 12, 'accent' if not dark else 'gray_dark', semi=True)
    if title:
        tx(s, title, M, .9, CW, 1.0, title_size, fg, semi=True, line_spacing=1.0)
    if source:
        tx(s, source, M, 7.02, CW - 1, .25, 10, sec)
    tx(s, f'{len(P.slides)}', SW - M - .6, 7.02, .6, .25, 10, sec, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = (f'{"BACKUP" if backup else "MAIN TALK"} | ~{secs} s\n\n' + notes_text.strip()
                                           + '\n\nSOURCE: ' + (source or ''))
    MANIFEST.append(dict(n=len(P.slides), title=title, act=act, source=source, seconds=secs, backup=backup))
    return s


def divider(num, title, items, secs=15):
    s = P.slides.add_slide(P.slide_layouts[6])
    _DARK[id(s)] = True
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb('bg_dark')
    set_fade(s, 700)
    tx(s, f'Act {num} of 6', M, 1.2, 6, .35, 14, 'gray_dark', semi=True)
    tx(s, f'{num:02}', M - .05, 1.5, 6, 2.6, 160, 'num_dark', semi=True)
    tx(s, title, M, 4.05, CW, 1.0, 48, 'white', semi=True)
    for i, it in enumerate(items):
        tx(s, it, M, 5.3 + i * .48, CW, .48, 19, 'gray_dark')
    tx(s, f'{len(P.slides)}', SW - M - .6, 7.02, .6, .25, 10, 'gray_dark', align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = f'ACT {num}: {title}\n' + ' · '.join(items)
    MANIFEST.append(dict(n=len(P.slides), title=f'Act {num}: {title}', act='Divider', source='', seconds=secs, backup=False))
    return s


# ----------------------------------------------------------------------------- composites
def key(s, text, y=6.15, color=None, size=21):
    """Bottom key sentence: short accent rule + one semibold line."""
    dark = is_dark(s)
    rect(s, M, y, .55, .045, color or 'accent')
    tx(s, text, M, y + .17, CW, .6, size, color if color in ('redtext', 'green', 'orange') else ('white' if dark else 'black'),
       semi=True)


def card(s, title, body, x, y, w, h, title_color=None, body_color=None, size=18, title_size=22, fill=None):
    dark = is_dark(s)
    rect(s, x, y, w, h, fill or ('fill_dark' if dark else 'fill'), radius=.22)
    tx(s, title, x + .35, y + .32, w - .7, .6, title_size, title_color or ('white' if dark else 'black'), semi=True)
    if body:
        tx(s, body, x + .35, y + 1.0, w - .7, h - 1.2, size, body_color or ('gray_dark' if dark else 'gray'),
           space_after=6, line_spacing=1.1)


def columns(s, items, y=TOP + .2, h=3.6, gap=.5, size=18, title_size=24, numbered=False, title_colors=None):
    """Typographic columns: [(title, body)]. No boxes; a hairline above each column."""
    dark = is_dark(s)
    n = len(items)
    w = (CW - gap * (n - 1)) / n
    for i, (t, b) in enumerate(items):
        x = M + i * (w + gap)
        hline(s, x, y, w, 'hair_dark' if dark else 'hair')
        yy = y + .3
        if numbered:
            tx(s, f'{i + 1:02}', x, yy, w, .3, 13, 'accent', semi=True); yy += .4
        tc = (title_colors or [None] * n)[i] or ('white' if dark else 'black')
        tx(s, t, x, yy, w, .7, title_size, tc, semi=True)
        tx(s, b, x, yy + .8, w, h - 1.3, size, 'gray_dark' if dark else 'gray', space_after=6, line_spacing=1.12)


def flow(s, items, y=TOP + .35, h=2.4, size=17, title_size=21, last_color=None):
    """Horizontal chain of light cards joined by thin arrows."""
    n = len(items)
    gap = .55
    w = (CW - gap * (n - 1)) / n
    for i, (t, b) in enumerate(items):
        x = M + i * (w + gap)
        rect(s, x, y, w, h, 'fill', radius=.22)
        tc = last_color if (i == n - 1 and last_color) else 'black'
        tx(s, t, x + .3, y + .3, w - .6, .6, title_size, tc, semi=True)
        tx(s, b, x + .3, y + .95, w - .6, h - 1.1, size, 'gray', line_spacing=1.1)
        if i < n - 1:
            arrow(s, x + w + .1, y + h / 2, x + w + gap - .1, y + h / 2)


def stat(s, number, label, detail, x, y, w, color='black', num_size=60):
    dark = is_dark(s)
    tx(s, number, x, y, w, 1.0, num_size, 'white' if (dark and color == 'black') else color, semi=True)
    tx(s, label, x, y + 1.1, w, .45, 21, 'white' if dark else 'black', semi=True)
    tx(s, detail, x, y + 1.55, w, .4, 16, 'gray_dark' if dark else 'gray')


def stats(s, cards, y=TOP + .3, cols=2, rowgap=2.3):
    n = cols
    gap = .6
    w = (CW - gap * (n - 1)) / n
    for i, (num, lab, det, col) in enumerate(cards):
        r, c = divmod(i, n)
        stat(s, num, lab, det, M + c * (w + gap), y + r * rowgap, w, col)


def table(s, headers, data, x=M, y=TOP, w=CW, colfracs=None, rowh=.6, size=17, highlight=(), hl_color='accent',
          first_semi=True, muted_cols=(), head_size=12):
    """Hairline table: small-caps grey header, thin rules, no fills."""
    dark = is_dark(s)
    n = len(headers)
    colfracs = colfracs or [1 / n] * n
    xs = [x]
    for f in colfracs[:-1]:
        xs.append(xs[-1] + w * f)
    ws = [w * f for f in colfracs]
    fg = 'white' if dark else 'black'
    for c, h in enumerate(headers):
        tx(s, h.upper(), xs[c] + .05, y, ws[c] - .1, .3, head_size, 'gray3', semi=True)
    hline(s, x, y + .4, w, fg if not dark else 'gray_dark', 1.0)
    hls = set(highlight if isinstance(highlight, (list, tuple, set)) else [highlight])
    for r, row in enumerate(data):
        yy = y + .4 + r * rowh
        for c, val in enumerate(row):
            col = hl_color if r in hls else (('gray_dark' if dark else 'gray') if c in muted_cols else fg)
            tx(s, str(val), xs[c] + .05, yy + (rowh - .4) / 2, ws[c] - .1, .4, size, col,
               semi=(c == 0 and first_semi) or r in hls, anchor=MSO_ANCHOR.MIDDLE)
        hline(s, x, yy + rowh, w, 'hair_dark' if dark else 'hair')


def bullets(s, items, y=TOP + .3, gap=1.25, size=24, body_size=18, x=M, w=CW):
    dark = is_dark(s)
    for i, (lead, body) in enumerate(items):
        yy = y + i * gap
        rect(s, x, yy + .15, .12, .12, 'accent', shape=MSO_SHAPE.OVAL)
        tx(s, lead, x + .4, yy, w - .4, .5, size, 'white' if dark else 'black', semi=True)
        if body:
            tx(s, body, x + .4, yy + .5, w - .4, gap - .5, body_size, 'gray_dark' if dark else 'gray')


def _asset(name):
    for d in ASSET_DIRS:
        p = d / (name + '.png')
        if p.exists():
            return p
    raise FileNotFoundError(name)


def picture(s, name, x, y, w, h):
    from PIL import Image
    path = _asset(name)
    im = Image.open(path); iw, ih = im.size
    scale = min(w / iw, h / ih); nw = iw * scale; nh = ih * scale
    pic = s.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2), width=Inches(nw), height=Inches(nh))
    pic.name = uid('pic')
    return pic


def steps(s, items, y=TOP + .3, h=3.2, colors=None):
    """Numbered process: big accent numeral, title, body — typographic, no boxes."""
    n = len(items)
    gap = .6
    w = (CW - gap * (n - 1)) / n
    body_h = max(.4, h - 1.7)
    for i, (t, b) in enumerate(items):
        x = M + i * (w + gap)
        tx(s, f'{i + 1}', x, y, w, .9, 48, (colors or ['accent'] * n)[i], semi=True)
        tx(s, t, x, y + 1.0, w, .6, 26, 'black', semi=True)
        tx(s, b, x, y + 1.7, w, body_h, 18, 'gray', line_spacing=1.12)
        if i < n - 1:
            tx(s, '→', x + w + .12, y + 1.0, .4, .6, 24, 'hair')


# ----------------------------------------------------------------------------- save
def save(pptx_path):
    P.save(pptx_path)
    (Path(pptx_path).parent / 'slide_manifest.json').write_text(json.dumps(MANIFEST, ensure_ascii=False, indent=2), encoding='utf-8')
