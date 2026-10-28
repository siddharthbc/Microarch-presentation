"""Reuse original paper figures and editable listing excerpts in the v2 deck.

Called by build.py, or patch an existing PPTX with --input/--output.
Only named slide bodies are replaced; the slide sequence and speaker notes survive.
"""
from pathlib import Path
import argparse
import json
import re

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from kit import tx, rect, key, table, hline, rgb, M, CW, TOP

HERE = Path(__file__).resolve().parent
ASSETS = HERE / 'paper_assets'
FIG_PAGES = {1:3, 2:4, 3:4, 5:6, 7:8, 8:8}
LISTINGS = {x['number']:x for x in json.loads((ASSETS / 'code_listings.json').read_text(encoding='utf-8'))['listings']}
TARGETS = {
  7:('A DSA operation is a descriptor-driven pipeline', 'Figure 1'),
  10:('Two channels cross two different boundaries', 'Figure 7'),
  14:('Clue #1', 'Figure 3'),
  15:('The reveal', 'Figure 2'),
  16:('Proof of', 'Listing 2'),
  17:('The drawers are independent', 'Listings 3 and 4'),
  19:('The money shot', 'Figure 5'),
  21:('Weapon #1', 'Listing 1, lines 7–12'),
  24:('Learning to hold the queue hostage', 'Listing 5'),
  25:('The trap', 'Listing 6'),
  31:('Attack #1', 'Figure 8'),
}


def clear_body(slide):
    for shape in list(slide.shapes):
        y = shape.top / 914400
        if 1.90 <= y < 7.0:
            el = shape._element
            el.getparent().remove(el)


def source_footer(slide, value):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.top/914400 >= 7 and shape.left/914400 < 2:
            p = shape.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = value
                for r in list(p.runs)[1:]:
                    r._r.getparent().remove(r._r)
            else:
                p.text = value
            break


def figure(slide, number, x, y, w, h):
    path = ASSETS / f'paper_fig{number}.png'
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w/iw, h/ih)
    nw, nh = iw*scale, ih*scale
    pic = slide.shapes.add_picture(str(path), Inches(x+(w-nw)/2), Inches(y+(h-nh)/2), width=Inches(nw), height=Inches(nh))
    pic.name = f'Original paper Figure {number} — page {FIG_PAGES[number]}'
    return pic


def code_panel(slide, number, x, y, w, h, size=17, timing=False):
    item = LISTINGS[number]
    code = item['timing_excerpt']['code'] if timing else item['code']
    label = f'LISTING {number}' + (' · LINES 7–12' if timing else '')
    rect(slide, x, y, w, h, 'fill', radius=.12)
    tx(slide, label+' · PAPER EXCERPT', x+.25, y+.16, w-.5, .27, 11, 'accent', semi=True)
    # Fit using the actual installed monospaced font; preserve the code's line breaks.
    lines = code.split('\n')
    for _ in range(15):
        font = ImageFont.truetype('C:/Windows/Fonts/consola.ttf', round(size*4))
        max_width_pt = max(font.getlength(line) for line in lines)/4
        if max_width_pt <= (w-.55)*72 and len(lines)*size*1.24 <= (h-.64)*72:
            break
        size -= .25
    if size < 13.5:
        raise ValueError(f'Listing {number} would be too small: {size} pt')
    box = slide.shapes.add_textbox(Inches(x+.25), Inches(y+.55), Inches(w-.5), Inches(h-.65))
    box.name = f'Editable paper Listing {number}' + (' lines 7-12' if timing else '')
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = Pt(size*1.24)
        # Preserve characters verbatim while rendering comments in a quieter color.
        for text_part, color in ((line.split('//',1)[0], 'black'), ('//'+line.split('//',1)[1], 'gray')) if '//' in line else ((line, 'black'),):
            if text_part:
                r = p.add_run(); r.text = text_part
                r.font.name = 'Consolas'; r.font.size = Pt(size); r.font.color.rgb = rgb(color)
    return code


def aside(slide, title, body, x, y, w=3.4, size=17):
    tx(slide, title, x, y, w, .38, 16, 'accent', semi=True)
    hline(slide, x, y+.42, w)
    tx(slide, body, x, y+.61, w, 2.9, size, 'black', line_spacing=1.12, space_after=0)


def apply(presentation):
    if len(presentation.slides) != 60:
        raise ValueError('Expected the current 60-slide v2 deck; refusing to patch a different version.')
    changes=[]
    for n,(title_fragment,asset) in TARGETS.items():
        s = presentation.slides[n-1]
        texts='\n'.join(sh.text for sh in s.shapes if sh.has_text_frame)
        if title_fragment not in texts:
            raise ValueError(f'Slide {n} does not match {title_fragment!r}')
        clear_body(s)
        extra=''
        if n == 7:
            figure(s,1,M,TOP+.05,CW,3.75)
            tx(s,'ORIGINAL FIGURE 1  ·  CPU → portal → work queue → engine → completion',M,5.84,CW,.25,12,'gray')
            key(s,'Follow one descriptor through the engine and its memory-translation path.',y=6.18,size=20)
            source='Paper Fig. 1, §II-A/B · p. 3 · original diagram; caption omitted'
            extra='Walk the original Figure 1 left to right: CPU socket and root complex, portal, work queues, batch fetcher and arbiter, processing unit, then the named Decode/Read/Compute/Write/Notify stages. ATS and physical memory requests are separate paths. These are functional stages, not a fixed cycle-by-cycle schedule. Use the stage names: the paper prose has inconsistent numeric stage references.'
        elif n == 10:
            figure(s,7,M,TOP+.3,7.55,3.55)
            aside(s,'Two sharing requirements','Red: DevTLB timing\nSame engine; queues may differ.\n\nTeal: SWQ contention\nSame shared work queue.',M+8.05,TOP+.25,3.45,17)
            key(s,'The shared resource determines which boundary the attacker can observe.',y=6.15,size=20)
            source='Paper Fig. 7, §§V–VI · p. 8 · original attack setup'
            extra='Figure 7 is the authors’ original cross-VM setup. Red traces denote the DevTLB path: separate queues can reach one engine. Teal traces denote same-SWQ contention. Sender/receiver describe a cooperative covert channel, while attacker/victim describe the side-channel case. Do not treat those role labels as four simultaneously required processes.'
        elif n == 14:
            figure(s,3,M,TOP+.2,7.05,3.8)
            tx(s,'SEMANTIC FIELD → SLOT',M+7.55,TOP+.1,3.95,.3,13,'accent',semi=True)
            table(s,['Operation','Fields touched'],[['NOOP','comp'],['Compare','src, src2, comp'],['Copy','src, dst, comp'],['Dualcast','src, dst, dst2, comp']],x=M+7.55,y=TOP+.62,w=3.95,colfracs=[.35,.65],rowh=.57,size=15,muted_cols=(1,))
            tx(s,'Completion writing is enabled\nin these examples.',M+7.55,5.55,3.95,.55,14,'gray')
            key(s,'The same bytes at offset 24 mean src2 for compare and dst for copy.',y=6.15,size=20)
            source='Paper Fig. 3, §IV-B · pp. 4–5 · original descriptor diagram'
            extra='The diagram is the actual 64-byte descriptor format from Figure 3. Explain offsets 8 (completion), 16 (source), 24 (dst/src2), and 40 (dst2). The operation table is a presenter annotation using semantic operation names. Listing 1 prints DSA_OPCODE_COMPVAL despite naming a two-source helper; do not present that printed constant as the encoding for a two-buffer compare.'
        elif n == 15:
            figure(s,2,M,TOP+.15,7.55,3.95)
            aside(s,'Read the original structure','One engine owns five slots:\nsrc · src2 · dst · dst2 · comp\n\nOne translation per slot.\n\nA new page in the same field\nreplaces the previous entry.',M+8.05,TOP+.15,3.45,17)
            key(s,'The slots cache translations; shared replacement capacity exposes activity.',y=6.15,size=20)
            source='Paper Fig. 2, §IV-A/B · p. 4 · original reverse-engineered structure'
            extra='Use the colored stacks at the top of Figure 2 to connect per-engine state to the five semantic fields. The lower diagram shows queue admission, batch fetching, arbitration, and processing. Five one-entry field slots are the authors’ observations on their tested platforms. Shared replacement capacity does not authorize using another process’s address translation.'
        elif n == 16:
            code_panel(s,2,M,TOP,CW,1.8,size=20)
            table(s,['Offset from base','Final access to base','What it establishes'],[['Within the same 4 KiB page','Hit','Page-granular reuse'],['Another page','Miss','One conflicting page replaces base']],y=3.99,colfracs=[.30,.24,.46],rowh=.62,size=18,highlight=1,hl_color='redtext',muted_cols=(2,))
            key(s,'For NOOP, base is the attacker’s completion-record page.',y=6.15,size=20)
            source='Paper Listing 2, §IV-B · p. 5 · verbatim code; helper setup omitted'
            extra='Read each of the three calls aloud: establish base; attempt a conflict; revisit base. For probe_noop the address names completion memory, not a payload source. The helper and required completion-record flags are omitted in the paper. Hit/miss observations follow the experiment; these three calls are not a complete executable attack.'
        elif n == 17:
            code_panel(s,3,M,TOP,5.55,2.65,size=15.5)
            code_panel(s,4,M+6.0,TOP,5.55,2.65,size=16.5)
            tx(s,'Change src; keep dst.\nThe destination translation survives.',M,4.93,5.55,.72,18,'black',semi=True,line_spacing=1.15)
            tx(s,'Same page / encoding position; different role.\nOnly src contributes the reuse hit.',M+6.0,4.93,5.55,.72,18,'black',semi=True,line_spacing=1.15)
            key(s,'Different semantic fields use different DevTLB slots.',y=6.15,size=20)
            source='Paper Listings 3–4, §IV-B · p. 5 · verbatim code; setup omitted'
            extra='Listing 3 changes the source page while keeping the destination, demonstrating independent source and destination state. Listing 4 places src2 and dst on the same page but uses different operations: src is reused, while the src2 translation does not provide a dst-slot hit. Preserve the original helper names and comments; do not silently repair their spelling.'
        elif n == 19:
            figure(s,5,M,TOP,6.0,4.45)
            table(s,['Case','Shared resource','Probe'],[['E0','same WQ + engine','Miss'],['E1','different WQs; same engine','Miss'],['E2','different engines','Hit']],x=M+6.4,y=TOP+.5,w=5.15,colfracs=[.15,.64,.21],rowh=.8,size=17,highlight=2,hl_color='green',muted_cols=())
            tx(s,'E1: separate queues still share DevTLB capacity.\nE2: separate engines preserve the primed entry.',M+6.4,5.32,5.15,.85,15,'gray',line_spacing=1.1,space_after=5)
            key(s,'In these experiments, eviction followed the shared engine.',y=6.35,size=20)
            source='Paper Fig. 5, §IV-B · p. 6 · original E0/E1/E2 diagram'
            extra='The original Figure 5 contains both the victim submission sketch and the attacker prime/wait/probe sketch. Match each diagram row to its result. E2 is a negative control for this tested engine boundary, not proof that every possible DSA configuration is secure or that sharing alone guarantees every probe observes activity.'
        elif n == 21:
            tx(s,'PRIME: own completion page   →   IDLE: allow victim activity   →   PROBE: time own request',M,TOP,CW,.52,17,'black',semi=True)
            code_panel(s,1,M,2.7,7.45,2.75,size=20,timing=True)
            aside(s,'What the timer measures','ENQCMD returns first.\n\nRDTSC brackets completion polling.\n\nThe result is the attacker’s own\ncompletion duration.',M+7.95,2.75,3.6,16)
            tx(s,'Listing 1 times a compare helper; §V-B uses NOOP for the attack. The timing pattern carries over.',M,5.72,CW,.5,15,'gray')
            key(s,'Slow completion suggests replacement; it does not reveal the victim’s address.',y=6.25,size=19)
            source='Paper Listing 1, lines 7–12 · p. 4; NOOP attack in §V-B · p. 8'
            extra='This is an exact excerpt from Listing 1, not a silently rewritten NOOP implementation. The paper’s complete Listing 1 prints COMPVAL in a helper described as comparing two buffers; those inconsistent setup lines are deliberately outside this excerpt. The source omits initialization, flags, and alignment. Explain that the timer begins after enqcmd and brackets completion polling. Section V-B applies this timing approach to completion-enabled NOOPs. A slow sample is an activity indicator subject to noise, not certain attribution to one victim.'
        elif n == 24:
            code_panel(s,5,M,TOP,7.55,2.75,size=17)
            tx(s,'Repeat the experiment after swapping\nthe two ENQCMD submission positions.',M,5.03,7.55,.7,19,'black',semi=True)
            aside(s,'What the experiment found','Work completion latency barely\nchanged when order was swapped.\n\nMeasured preference:\nordinary work over batch work.\n\nA batch is an unreliable anchor.',M+8.0,TOP+.1,3.55,16)
            key(s,'Arbiter priority motivates a long ordinary copy for congestion.',y=6.15,size=20)
            source='Paper Listing 5, §IV-C · p. 6 · verbatim experiment; setup omitted'
            extra='Listing 5 is an arbitration experiment, not the congestion loop. The authors swap submission order and compare ordinary-work completion latency. The results support priority for ordinary work over batch-buffer descriptors in the tested implementation. The explicit SFENCE separates the two submissions in this microbenchmark. Queue admission checks are a different part of the design, already visible in Figure 2.'
        elif n == 25:
            code_panel(s,6,M,TOP,7.55,3.75,size=16)
            aside(s,'Connect code to queue state','mass_desc = long ordinary copy.\n\nwq_size − 2 additional descriptors.\n\nTarget: N − 1 occupied positions.\n\nprobe() reads acceptance / retry,\nnot the operation’s completion.',M+8.0,TOP+.05,3.55,16)
            tx(s,'Paper observation: congestion became visible for copies ≥ 32 MiB. Descriptor preparation is omitted.',M,5.93,CW,.4,14,'gray')
            key(s,'Keep the long copy active while leaving one admission position available.',y=6.35,size=19)
            source='Paper Listing 6, §§IV-C/V-C · pp. 7–8 · verbatim code; setup omitted'
            extra='The paper submits one long memcpy and then wq_size−2 lightweight descriptors, aiming at a near-full queue during the observation window. A victim may use the remaining admission capacity. The probe helper’s nonzero result corresponds to retry/full-queue feedback in this experiment; it does not measure completion or identify a victim with certainty. Descriptor preparation, initial completion state, and retry policy are omitted from this pedagogical listing. The idle interval must be shorter than the anchor copy’s execution.'
        elif n == 31:
            figure(s,8,M,TOP+.05,CW,3.65)
            tx(s,'Website traffic → DSA-accelerated memif transfers → shared-engine activity → miss-count trace',M,5.8,CW,.5,17,'black',semi=True,align=PP_ALIGN.CENTER)
            key(s,'The spy observes its own probes; the classifier supplies the website label.',y=6.3,size=20)
            source='Paper Fig. 8, §VI-B · pp. 8–10 · original VPP/memif setup'
            extra='The original Figure 8 locates the DSA spy in VM2 and the victim applications in VM1. VPP/memif uses DSA for packet movement, so workload bursts can affect the shared engine. The attacker monitors its own completion latency and aggregates inferred misses. It does not read packet contents or directly discover a browser process; the evaluated workload is controlled browsing, with a classifier trained on known website labels.'
        source_footer(s,source)
        old=s.notes_slide.notes_text_frame.text
        # Re-applying this patch updates the note supplement instead of duplicating it.
        old=old.split('\n\nPAPER VISUAL / CODE WALKTHROUGH:')[0]
        s.notes_slide.notes_text_frame.text=old+'\n\nPAPER VISUAL / CODE WALKTHROUGH:\n'+extra+'\n\nUPDATED SOURCE: '+source
        changes.append({'slide':n,'asset':asset,'source':source})
    return changes


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--input',required=True,type=Path)
    parser.add_argument('--output',required=True,type=Path)
    args=parser.parse_args()
    p=Presentation(args.input)
    changes=apply(p)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    p.save(args.output)
    (args.output.parent/'paper_visuals_manifest.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'slides':len(p.slides),'updated_slides':[c['slide'] for c in changes],'output':str(args.output)},indent=2))


if __name__=='__main__':
    main()
